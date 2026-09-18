import io
from pathlib import Path
from xml.sax.saxutils import escape
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from .schema import FIELDS

def safe_cell(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@')):return "'"+value
    return value

def simulation_rows(sim):
    return [
        ['Показатель','Значение'],
        ['Критерий пропускной способности','Выполнен' if sim['passed'] else 'Не выполнен: экономический эффект условен'],
        ['Горизонт, ч',sim['hours']],['Seed',sim['seed']],
        ['Целевой поток, операций/ч',round(sim['expected_throughput'],2)],
        ['Фактический поток, операций/ч',round(sim['throughput'],2)],
        ['Выполнено операций',sim['completed']],['Осталось в очереди',sim['queued']],
        ['В работе на конец',sim['in_progress']],['Ожидание начала задачи p95, мин',round(sim['p95_wait_minutes'],2)],
        ['Рабочая загрузка, %',round(sim['utilization'],2)],
        *[['Ограничение',x] for x in sim['bottlenecks']],
        *[['Допущение',x] for x in sim['notes']],
    ]

def excel_report(result):
    w=Workbook();s=w.active;s.title='Сравнение'
    s.append(['Робоскоп · оценка роботизации склада'])
    s.append(['Версия модели',result['model_version'],'Дата',result['calculated_at']])
    s.append(['Предварительная оценка. Требует проверки при обследовании объекта.'])
    s.append(['Сценарий','Первоначальные вложения, ₽','Годовые затраты, ₽','Годовой эффект, ₽','Окупаемость, лет','ROI по ТЗ, %','TCO, ₽','Чистый эффект, ₽'])
    for sc in result['scenarios']:s.append([sc['label']]+[sc[k] for k in ['capex','opex','annual_effect','payback','roi','tco','net_effect']])
    p=w.create_sheet('Параметры');p.append(['Параметр','Значение','Единица','Источник / допущение'])
    for f in FIELDS:p.append([f['label'],result['inputs'][f['key']],f['unit'],safe_cell(result.get('assumptions',{}).get(f['key'],{}).get('source',f['source']))])
    p.append(['Модель оборудования',safe_cell(result['product']['name'])]);p.append(['Роботы',result['fleet']['robots']]);p.append(['Зарядки',result['fleet']['chargers']])
    flow=w.create_sheet('Денежный поток');flow.append(['Сценарий','Год','Годовые затраты, ₽','Замена АКБ, ₽','Эффект за год, ₽','Накопленный чистый эффект, ₽'])
    for sc in result['scenarios']:
        for y in sc['cashflow']:flow.append([sc['label']]+[y[k] for k in ['year','opex','replacement','effect','cumulative']])
    detail=w.create_sheet('Статьи затрат');detail.append(['Сценарий','Тип','Статья','Сумма, ₽'])
    for sc in result['scenarios']:
        for group in ['capex_items','opex_items']:
            for name,value in sc[group].items():detail.append([sc['label'],'CAPEX' if group=='capex_items' else 'OPEX',name,value])
    sens=w.create_sheet('Чувствительность');sens.append(['Параметр','Изменение, %','Сценарий','Роботы','TCO, ₽','Эффект, ₽/год','Окупаемость, лет'])
    for row in result['sensitivity']:
        for sc in row['scenarios']:sens.append([row['label'],row['delta'],sc['key'],row['robots'],sc['tco'],sc['annual_effect'],sc['payback']])
    formulas=w.create_sheet('Формулы и ограничения');formulas.append(['Показатель','Формула','Примечание'])
    for row in result['formulas']:formulas.append(row)
    for warning in result['warnings']:formulas.append(['Ограничение',warning])
    for source in result['product']['sources']:formulas.append(['Источник',safe_cell(source['title']),source.get('date','')+' · '+source.get('status','')])
    if result.get('simulation_summary'):
        sim=w.create_sheet('Симуляция')
        for row in simulation_rows(result['simulation_summary']):sim.append(row)
    for sheet in w:
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='14273D');cell.font=Font(color='FFFFFF',bold=True)
        for col in sheet.columns:
            letter=col[0].column_letter;sheet.column_dimensions[letter].width=22 if letter!='A' else 40
            for cell in col:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,(int,float)):cell.number_format='#,##0.00'
        for row in sheet:sheet.row_dimensions[row[0].row].height=32
    p.column_dimensions['D'].width=85;formulas.column_dimensions['B'].width=85;formulas.column_dimensions['C'].width=60
    for row in formulas:formulas.row_dimensions[row[0].row].height=48
    out=io.BytesIO();w.save(out);return out.getvalue()

def pdf_report(result):
    font=settings.BASE_DIR/'static/fonts/DejaVuSans.ttf'
    bold=settings.BASE_DIR/'static/fonts/DejaVuSans-Bold.ttf'
    if 'DejaVu' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('DejaVu',str(font)));pdfmetrics.registerFont(TTFont('DejaVuBold',str(bold)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Ru',fontName='DejaVu',fontSize=9,leading=14,spaceAfter=8,textColor=colors.HexColor('#25364a')))
    styles.add(ParagraphStyle(name='RuTitle',fontName='DejaVuBold',fontSize=23,leading=29,spaceAfter=20))
    styles.add(ParagraphStyle(name='RuH',fontName='DejaVuBold',fontSize=13,leading=18,spaceBefore=14,spaceAfter=10))
    def p(text,style='Ru'):return Paragraph(escape(str(text)),styles[style])
    def table(rows,widths):
        t=Table([[p('—' if v is None else v) for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf4')),('VALIGN',(0,0),(-1,-1),'TOP'),
                              ('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#adbac8')),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#dde4ed')),
                              ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        return t
    def money(v):return f'{v:,.0f}'.replace(',',' ')+' ₽'
    out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=(595.28,841.89),rightMargin=40,leftMargin=40,topMargin=40,bottomMargin=40)
    story=[p('Робоскоп','RuTitle'),p('Предварительная оценка роботизации склада','RuH'),
           p('Дата: '+result['calculated_at'][:10]+' · Модель: '+result['model_version']),
           p('Решение: '+result['product']['name']),p(f"Оборудование: {result['fleet']['robots']} роботов, {result['fleet']['chargers']} зарядных станций."),
           p('Результат является предварительной оценкой и требует верификации при обследовании объекта. Все суммы в номинальных рублях, в единой ценовой базе. Налоги, дисконтирование и финансирование отдельно не моделируются.'),
           p('Сравнение сценариев','RuH')]
    rows=[['Показатель']+[s['label'] for s in result['scenarios']]]
    for label,key in [('Первоначальные вложения','capex'),('Годовые затраты','opex'),('Годовой эффект','annual_effect'),('TCO за горизонт','tco'),('Чистый эффект за горизонт','net_effect')]:
        rows.append([label]+[money(s[key]) for s in result['scenarios']])
    rows.append(['Окупаемость, лет']+[f"{s['payback']:.2f}" if s['payback'] is not None else 'Не применимо' for s in result['scenarios']])
    rows.append(['ROI по ТЗ, %']+[f"{s['roi']:.1f}" if s['roi'] is not None else 'Не применимо' for s in result['scenarios']])
    story.append(table(rows,[158,119,119,119]))
    if result.get('simulation_summary'):
        story.append(p('Проверка производительности','RuH'))
        story.append(table(simulation_rows(result['simulation_summary']),[285,230]))
    else:
        story.append(p('Симуляция не приложена. Достижимость рассчитанной производительности не проверена.'))
    story.append(p('Ограничения и риски','RuH'))
    story.extend(p('• '+x) for x in result['warnings'])
    story.append(PageBreak());story.append(p('Исходные параметры и допущения','RuH'))
    story.append(table([['Показатель','Значение','Источник / допущение']]+[[f['label'],str(result['inputs'][f['key']])+' '+f['unit'],result.get('assumptions',{}).get(f['key'],{}).get('source',f['source'])] for f in FIELDS],[168,92,255]))
    for sc in result['scenarios'][1:]:
        story.append(p(sc['label']+' · состав затрат','RuH'))
        story.append(table([['Статья','Сумма']]+[[k,money(v)] for k,v in sc['capex_items'].items()]+[['Ежегодно: '+k,money(v)] for k,v in sc['opex_items'].items()],[365,150]))
        story.append(p('Денежный поток','RuH'))
        story.append(table([['Год','Эффект за год','Замена АКБ','Накопленный чистый эффект']]+[[y['year'],money(y['effect']),money(y['replacement']),money(y['cumulative'])] for y in sc['cashflow']],[40,155,140,180]))
    story.append(p('Формулы','RuH'));story.append(table([['Показатель','Зависимость и пояснение']]+[[r[0],r[1]+'. '+r[2]] for r in result['formulas']],[145,370]))
    story.append(p('Источники данных','RuH'))
    for source in result['product']['sources']:story.append(p(source['title']+' · '+source.get('date','')+' · '+source.get('status','')+' '+source.get('url','')))
    def footer(canvas,doc):
        canvas.setFont('DejaVu',8);canvas.setFillColor(colors.HexColor('#748095'));canvas.drawString(40,23,'Робоскоп · Предварительная оценка');canvas.drawRightString(555,23,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()
